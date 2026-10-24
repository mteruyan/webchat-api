"""
reportes.py — reporte de conversaciones por asesor para auditoría de calidad (Excel).

Genera un .xlsx con tres hojas:
  1. Resumen        — indicadores por asesor (cantidad, tiempos de respuesta, duración...).
  2. Conversaciones — una fila por conversación con todos los datos y tiempos.
  3. Transcripción  — cada mensaje (fecha/hora, quién, texto) para revisar el trato.

Reglas de fechas: todo en hora de Lima. El filtro es por la fecha de ASIGNACIÓN al
asesor (o la de inicio si la conversación no llegó a un asesor).
"""

import io
import json
import re
from datetime import datetime, timedelta, timezone, date

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from db import get_connection, release_connection

LIMA_TZ = timezone(timedelta(hours=-5))
MAX_CONVERSACIONES = 20000

AZUL = "0A4C8C"
FMT_FECHA_HORA = "dd/mm/yyyy hh:mm:ss"
FMT_DURACION = "[h]:mm:ss"

_ILEGALES = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


# ---------------------------------------------------------------------------
# Rangos de fecha
# ---------------------------------------------------------------------------
def resolver_rango(rango: str, desde: str = None, hasta: str = None):
    """Devuelve (fecha_desde, fecha_hasta, etiqueta), ambas inclusivas, en hora de Lima."""
    hoy = datetime.now(LIMA_TZ).date()
    rango = (rango or "hoy").lower()
    if rango == "hoy":
        return hoy, hoy, "Hoy"
    if rango == "ayer":
        d = hoy - timedelta(days=1)
        return d, d, "Ayer"
    if rango == "semana":
        return hoy - timedelta(days=6), hoy, "Última semana (7 días)"
    if rango == "mes_actual":
        return hoy.replace(day=1), hoy, "Mes actual"
    if rango == "mes_pasado":
        primero_este = hoy.replace(day=1)
        ultimo_pasado = primero_este - timedelta(days=1)
        return ultimo_pasado.replace(day=1), ultimo_pasado, "Mes pasado"
    if rango == "personalizado":
        try:
            d1 = datetime.strptime(desde or "", "%Y-%m-%d").date()
            d2 = datetime.strptime(hasta or "", "%Y-%m-%d").date()
        except ValueError:
            raise ValueError("Las fechas del rango personalizado no son válidas.")
        if d2 < d1:
            raise ValueError("La fecha 'hasta' no puede ser anterior a 'desde'.")
        if (d2 - d1).days > 366:
            raise ValueError("El rango no puede superar un año.")
        return d1, d2, "Personalizado"
    raise ValueError("Rango no reconocido.")


# ---------------------------------------------------------------------------
# Utilidades de celdas
# ---------------------------------------------------------------------------
def _naive_lima(dt):
    """Excel no admite zonas horarias: se convierte a hora de Lima sin tz."""
    if dt is None:
        return None
    if dt.tzinfo is not None:
        dt = dt.astimezone(LIMA_TZ).replace(tzinfo=None)
    return dt


def _texto(valor, limite=32000):
    if valor is None:
        return None
    s = _ILEGALES.sub("", str(valor))
    return s[:limite]


def _celda(ws, valor, formato=None, negrita=False, color=None, fondo=None, ajustar=False):
    c = WriteOnlyCell(ws, value=valor)
    if isinstance(valor, str) and valor[:1] in ("=", "+", "-", "@"):
        c.data_type = "s"  # un texto que empieza con "=" no debe ejecutarse como fórmula
    if formato:
        c.number_format = formato
    if negrita or color:
        c.font = Font(bold=negrita, color=color)
    if fondo:
        c.fill = PatternFill("solid", fgColor=fondo)
    if ajustar:
        c.alignment = Alignment(wrap_text=True, vertical="top")
    return c


def _encabezado(ws, titulos):
    return [_celda(ws, t, negrita=True, color="FFFFFF", fondo=AZUL, ajustar=True) for t in titulos]


def _promedio(valores):
    valores = [v for v in valores if v is not None]
    return (sum(valores, timedelta()) / len(valores)) if valores else None


def _suma(valores):
    valores = [v for v in valores if v is not None]
    return sum(valores, timedelta()) if valores else None


def _delta(a, b):
    """b - a si ambos existen y no es negativo."""
    if a is None or b is None:
        return None
    d = b - a
    return d if d >= timedelta() else None


def _parsear_json(valor):
    if isinstance(valor, (dict, list)):
        return valor
    if isinstance(valor, str):
        try:
            return json.loads(valor)
        except Exception:
            return None
    return None


# ---------------------------------------------------------------------------
# Consultas
# ---------------------------------------------------------------------------
def _columnas(cur, tabla):
    cur.execute(
        "SELECT column_name FROM information_schema.columns WHERE table_name = %s", (tabla,)
    )
    return {r[0] for r in cur.fetchall()}


def _texto_mensaje(emisor, tipo, contenido):
    """Texto legible del mensaje según su tipo (para la hoja de transcripción)."""
    if tipo == "archivo":
        nombre = (contenido or "").rsplit("/", 1)[-1]
        return f"📎 Archivo: {nombre}" if nombre else "📎 Archivo"
    if tipo == "opciones":
        datos = _parsear_json(contenido)
        etiquetas = []
        lista = datos if isinstance(datos, list) else (datos or {}).get("opciones") if isinstance(datos, dict) else None
        for o in (lista or []):
            if isinstance(o, dict):
                etiquetas.append(str(o.get("label") or o.get("texto") or o.get("id") or ""))
        return "[Menú] " + " | ".join(e for e in etiquetas if e) if etiquetas else "[Menú del bot]"
    if tipo == "boton_click":
        return f"[Botón] {contenido or ''}"
    if tipo == "formulario":
        datos = _parsear_json(contenido)
        if emisor == "usuario" and isinstance(datos, dict):
            return "[Formulario enviado] " + "; ".join(f"{k}: {v}" for k, v in datos.items())
        return "[Formulario enviado]" if emisor == "usuario" else "[Formulario del bot]"
    return contenido or ""


def generar_reporte(desde: date, hasta: date, etiqueta_rango: str, generado_por: str,
                    asesor_id=None, campana_id=None, incluir_solo_bot=False,
                    incluir_transcripcion=True, campana_ids=None) -> bytes:
    inicio = datetime.combine(desde, datetime.min.time(), tzinfo=LIMA_TZ)
    fin = datetime.combine(hasta + timedelta(days=1), datetime.min.time(), tzinfo=LIMA_TZ)

    conn = get_connection()
    try:
        with conn.cursor() as cur:
            cols_c = _columnas(cur, "conversaciones")
            cols_u = _columnas(cur, "usuarios_webchat")

            def opc(col, tabla="c", cols=cols_c):
                return f"{tabla}.{col}" if col in cols else "NULL"

            filtros = ["COALESCE(c.asignado_en, c.iniciado_en) >= %s",
                       "COALESCE(c.asignado_en, c.iniciado_en) < %s"]
            params = [inicio, fin]
            if not incluir_solo_bot:
                filtros.append("c.asesor_id IS NOT NULL")
            if asesor_id:
                filtros.append("c.asesor_id = %s")
                params.append(int(asesor_id))
            if campana_id:
                filtros.append("c.campana_id = %s")
                params.append(int(campana_id))
            if campana_ids is not None:  # alcance de un supervisor de campaña ([] = ninguna)
                filtros.append("c.campana_id = ANY(%s)")
                params.append([int(i) for i in campana_ids])

            cur.execute(
                f"""
                SELECT c.id, c.estado, c.iniciado_en, {opc('entrada_cola')}, c.asignado_en, {opc('cerrado_en')},
                       {opc('cerrado_por')}, {opc('bandera')}, {opc('origen')}, {opc('pagina_origen')},
                       cam.nombre, a.id, a.nombre_display, a.usuario,
                       u.nombre, u.telefono, u.dni, {opc('ruc', 'u', cols_u)},
                       et.nombre
                FROM conversaciones c
                JOIN usuarios_webchat u ON u.id = c.usuario_id
                LEFT JOIN campanas cam ON cam.id = c.campana_id
                LEFT JOIN asesores a ON a.id = c.asesor_id
                LEFT JOIN etiquetas_cierre et ON et.id = c.etiqueta_cierre_id
                WHERE {' AND '.join(filtros)}
                ORDER BY a.nombre_display NULLS LAST, COALESCE(c.asignado_en, c.iniciado_en), c.id
                LIMIT {MAX_CONVERSACIONES + 1}
                """,
                params,
            )
            filas = cur.fetchall()
            if len(filas) > MAX_CONVERSACIONES:
                raise ValueError(
                    f"El rango tiene más de {MAX_CONVERSACIONES:,} conversaciones. Reduce el rango o filtra por asesor/campaña.")

            ids = [f[0] for f in filas]
            mensajes_por_conv = {}
            nombres_asesores = {}
            intervenciones = {}
            if ids:
                cur.execute("SELECT id, nombre_display FROM asesores")
                nombres_asesores = {i: n for i, n in cur.fetchall()}
                cur.execute(
                    """SELECT conversacion_id, emisor_tipo, emisor_id, tipo_contenido, contenido,
                              enviado_en, leido_en, metadata
                       FROM mensajes WHERE conversacion_id = ANY(%s)
                       ORDER BY conversacion_id, enviado_en, id""",
                    (ids,),
                )
                for conv_id, *resto in cur.fetchall():
                    mensajes_por_conv.setdefault(conv_id, []).append(resto)
                try:
                    cur.execute("SAVEPOINT sp_interv")
                    cur.execute(
                        """SELECT conversacion_id, COUNT(*) FROM supervisor_intervenciones
                           WHERE conversacion_id = ANY(%s) GROUP BY conversacion_id""",
                        (ids,),
                    )
                    intervenciones = {c: n for c, n in cur.fetchall()}
                    cur.execute("RELEASE SAVEPOINT sp_interv")
                except Exception:
                    cur.execute("ROLLBACK TO SAVEPOINT sp_interv")
    finally:
        release_connection(conn)

    # ------------------------------------------------------------------
    # Métricas por conversación
    # ------------------------------------------------------------------
    registros = []
    for (cid, estado, iniciado, entrada_cola, asignado, cerrado, cerrado_por, bandera, origen, pagina,
         campana, a_id, a_nombre, a_usuario, c_nombre, c_tel, c_dni, c_ruc, etiqueta) in filas:
        iniciado, entrada_cola, asignado, cerrado = map(_naive_lima, (iniciado, entrada_cola, asignado, cerrado))
        msgs = mensajes_por_conv.get(cid, [])
        n_cli = n_ase = n_bot = n_sup = n_archivos = n_transf = 0
        pendiente = None
        tiempos_resp = []
        primera_resp = None
        for emisor, _eid, tipo, contenido, enviado, _leido, _meta in msgs:
            enviado = _naive_lima(enviado)
            if emisor == "usuario":
                n_cli += 1
                if tipo == "archivo":
                    n_archivos += 1
                if asignado and enviado and enviado >= asignado and tipo != "boton_click" and pendiente is None:
                    pendiente = enviado
            elif emisor == "asesor":
                n_ase += 1
                if asignado and enviado and enviado >= asignado:
                    if primera_resp is None:
                        primera_resp = enviado - asignado
                    if pendiente is not None:
                        tiempos_resp.append(max(enviado - pendiente, timedelta()))
                        pendiente = None
            elif emisor == "supervisor":
                n_sup += 1
            else:
                n_bot += 1
                if "Conversación transferida" in (contenido or ""):
                    n_transf += 1

        fin_atencion = cerrado
        registros.append({
            "id": cid, "estado": estado, "iniciado": iniciado, "entrada_cola": entrada_cola,
            "asignado": asignado, "cerrado": cerrado, "cerrado_por": cerrado_por, "bandera": bandera,
            "origen": origen, "pagina": pagina, "campana": campana, "asesor": a_nombre, "asesor_usuario": a_usuario,
            "cliente": c_nombre, "telefono": c_tel, "dni": c_dni, "ruc": c_ruc, "etiqueta": etiqueta,
            "t_bot": _delta(iniciado, entrada_cola or asignado),
            "t_cola": _delta(entrada_cola, asignado),
            "primera_resp": primera_resp,
            "resp_prom": _promedio(tiempos_resp),
            "resp_max": max(tiempos_resp) if tiempos_resp else None,
            "dur_total": _delta(iniciado, cerrado),
            "dur_atencion": _delta(asignado, fin_atencion),
            "n_cli": n_cli, "n_ase": n_ase, "n_bot": n_bot, "n_sup": n_sup,
            "interv": intervenciones.get(cid, 0), "transf": n_transf, "archivos": n_archivos,
            "sin_resp": "SÍ" if (a_id and asignado and n_ase == 0) else "NO",
        })

    # ------------------------------------------------------------------
    # Excel
    # ------------------------------------------------------------------
    wb = Workbook(write_only=True)

    # --- Hoja Resumen -------------------------------------------------
    ws = wb.create_sheet("Resumen")
    ws.column_dimensions["A"].width = 30
    for i in range(2, 13):
        ws.column_dimensions[get_column_letter(i)].width = 19
    ws.append([_celda(ws, "Reporte de conversaciones por asesor — Auditoría de calidad", negrita=True)])
    ws.append([_celda(ws, "Rango:"), _celda(ws, f"{etiqueta_rango}: {desde.strftime('%d/%m/%Y')} al {hasta.strftime('%d/%m/%Y')}")])
    ws.append([_celda(ws, "Filtro de fecha:"), _celda(ws, "Fecha de asignación al asesor (o de inicio si no llegó a un asesor). Hora de Lima.")])
    ws.append([_celda(ws, "Generado por:"), _celda(ws, _texto(generado_por))])
    ws.append([_celda(ws, "Generado el:"), _celda(ws, datetime.now(LIMA_TZ).replace(tzinfo=None), FMT_FECHA_HORA)])
    ws.append([_celda(ws, "Conversaciones:"), _celda(ws, len(registros))])
    ws.append([])
    ws.append(_encabezado(ws, [
        "Asesor", "Conversaciones", "Cerradas", "En curso", "Msgs del asesor",
        "Espera en cola (prom.)", "1ª respuesta (prom.)", "Tiempo de respuesta (prom.)",
        "Tiempo de respuesta (máx.)", "Duración de atención (prom.)", "Tiempo total de atención",
        "Sin respuesta del asesor"]))
    por_asesor = {}
    for r in registros:
        por_asesor.setdefault(r["asesor"] or "(Sin asesor — solo bot)", []).append(r)
    for nombre, rs in sorted(por_asesor.items()):
        resp_prom = _promedio([r["resp_prom"] for r in rs])
        resp_max = max([r["resp_max"] for r in rs if r["resp_max"]], default=None)
        ws.append([
            _celda(ws, _texto(nombre), negrita=True), _celda(ws, len(rs)),
            _celda(ws, sum(1 for r in rs if r["estado"] == "cerrada")),
            _celda(ws, sum(1 for r in rs if r["estado"] != "cerrada")),
            _celda(ws, sum(r["n_ase"] for r in rs)),
            _celda(ws, _promedio([r["t_cola"] for r in rs]), FMT_DURACION),
            _celda(ws, _promedio([r["primera_resp"] for r in rs]), FMT_DURACION),
            _celda(ws, resp_prom, FMT_DURACION), _celda(ws, resp_max, FMT_DURACION),
            _celda(ws, _promedio([r["dur_atencion"] for r in rs]), FMT_DURACION),
            _celda(ws, _suma([r["dur_atencion"] for r in rs]), FMT_DURACION),
            _celda(ws, sum(1 for r in rs if r["sin_resp"] == "SÍ")),
        ])
    ws.append([])
    ws.append([_celda(ws, "Notas:", negrita=True)])
    for nota in (
        "Espera en cola: desde que el bot derivó la conversación hasta que se asignó a un asesor.",
        "1ª respuesta: desde la asignación hasta el primer mensaje del asesor.",
        "Tiempo de respuesta: desde que el cliente escribe (estando asignado) hasta que el asesor le responde; se promedia por cada intercambio.",
        "Duración de atención: desde la asignación hasta el cierre. Duración total: desde que el cliente abrió el chat hasta el cierre. Las conversaciones aún abiertas no tienen duración.",
        "Si una conversación fue transferida, aparece con el último asesor que la atendió.",
        "Los tiempos se muestran como horas:minutos:segundos.",
    ):
        ws.append([_celda(ws, nota)])

    # --- Hoja Conversaciones ------------------------------------------
    ws = wb.create_sheet("Conversaciones")
    titulos = [
        ("ID conversación", 13), ("Fecha", 12), ("Hora de inicio", 12), ("Campaña", 24),
        ("Asesor", 24), ("Usuario del asesor", 16),
        ("Cliente", 26), ("Celular", 14), ("DNI", 12), ("RUC", 14),
        ("Estado", 12), ("Etiqueta de cierre", 26), ("Cerrado por", 13), ("Bandera", 10),
        ("Origen", 14), ("Página de origen", 20),
        ("Inicio del chat", 19), ("Entrada a cola", 19), ("Asignación al asesor", 19), ("Cierre", 19),
        ("Tiempo en el bot", 14), ("Espera en cola", 14), ("1ª respuesta del asesor", 15),
        ("Tiempo de respuesta (prom.)", 15), ("Tiempo de respuesta (máx.)", 15),
        ("Duración de atención", 14), ("Duración total", 14),
        ("Msgs del cliente", 10), ("Msgs del asesor", 10), ("Msgs del bot", 10), ("Msgs del supervisor", 11),
        ("Intervenciones de supervisor", 14), ("Transferencias", 12), ("Archivos del cliente", 11),
        ("Sin respuesta del asesor", 12),
    ]
    for i, (_, ancho) in enumerate(titulos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.freeze_panes = "A2"
    ws.append(_encabezado(ws, [t for t, _ in titulos]))
    for r in registros:
        ref = r["asignado"] or r["iniciado"]
        ws.append([
            _celda(ws, r["id"]),
            _celda(ws, ref.date() if ref else None, "dd/mm/yyyy"),
            _celda(ws, ref.time() if ref else None, "hh:mm:ss"),
            _celda(ws, _texto(r["campana"])), _celda(ws, _texto(r["asesor"])), _celda(ws, _texto(r["asesor_usuario"])),
            _celda(ws, _texto(r["cliente"])), _celda(ws, _texto(r["telefono"])), _celda(ws, _texto(r["dni"])),
            _celda(ws, _texto(r["ruc"])),
            _celda(ws, r["estado"]), _celda(ws, _texto(r["etiqueta"])), _celda(ws, r["cerrado_por"]),
            _celda(ws, r["bandera"]), _celda(ws, r["origen"]), _celda(ws, _texto(r["pagina"])),
            _celda(ws, r["iniciado"], FMT_FECHA_HORA), _celda(ws, r["entrada_cola"], FMT_FECHA_HORA),
            _celda(ws, r["asignado"], FMT_FECHA_HORA), _celda(ws, r["cerrado"], FMT_FECHA_HORA),
            _celda(ws, r["t_bot"], FMT_DURACION), _celda(ws, r["t_cola"], FMT_DURACION),
            _celda(ws, r["primera_resp"], FMT_DURACION), _celda(ws, r["resp_prom"], FMT_DURACION),
            _celda(ws, r["resp_max"], FMT_DURACION), _celda(ws, r["dur_atencion"], FMT_DURACION),
            _celda(ws, r["dur_total"], FMT_DURACION),
            _celda(ws, r["n_cli"]), _celda(ws, r["n_ase"]), _celda(ws, r["n_bot"]), _celda(ws, r["n_sup"]),
            _celda(ws, r["interv"]), _celda(ws, r["transf"]), _celda(ws, r["archivos"]),
            _celda(ws, r["sin_resp"]),
        ])
    ws.auto_filter.ref = f"A1:{get_column_letter(len(titulos))}{len(registros) + 1}"

    # --- Hoja Transcripción -------------------------------------------
    if incluir_transcripcion:
        ws = wb.create_sheet("Transcripción")
        for letra, ancho in zip("ABCDEFG", (13, 24, 19, 26, 90, 14, 28)):
            ws.column_dimensions[letra].width = ancho
        ws.freeze_panes = "A2"
        ws.append(_encabezado(ws, ["ID conversación", "Asesor", "Fecha y hora", "Quién escribe",
                                   "Mensaje", "Leído por el destinatario", "Observación"]))
        filas_t = 0
        for r in registros:
            for emisor, e_id, tipo, contenido, enviado, leido, meta in mensajes_por_conv.get(r["id"], []):
                meta = _parsear_json(meta) or {}
                if emisor == "usuario":
                    quien = "Cliente"
                elif emisor == "asesor":
                    quien = f"Asesor: {nombres_asesores.get(e_id) or r['asesor'] or ''}".strip()
                elif emisor == "supervisor":
                    quien = f"Supervisor: {nombres_asesores.get(e_id) or ''}".strip()
                else:
                    quien = "Bot (Sofi)"
                obs = []
                if meta.get("eliminado_en"):
                    por = meta.get("eliminado_por")
                    obs.append("Mensaje eliminado" + (f" por {nombres_asesores.get(por, por)}" if por else ""))
                if meta.get("editado_en"):
                    obs.append("Editado")
                ws.append([
                    _celda(ws, r["id"]), _celda(ws, _texto(r["asesor"])),
                    _celda(ws, _naive_lima(enviado), FMT_FECHA_HORA), _celda(ws, _texto(quien)),
                    _celda(ws, _texto(_texto_mensaje(emisor, tipo, contenido)), ajustar=True),
                    _celda(ws, "Sí" if leido else "No"), _celda(ws, "; ".join(obs) or None),
                ])
                filas_t += 1
        ws.auto_filter.ref = f"A1:G{filas_t + 1}"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
