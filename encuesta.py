"""encuesta.py — encuesta de satisfacción al cerrar una conversación atendida por un asesor.

Cuándo se aplica: el asesor (o un supervisor) cierra una conversación que SÍ tuvo asesor asignado y cuya
campaña tiene activada la encuesta (campanas.encuesta_satisfaccion; se activa en Supervisor > Campañas o con
migracion_encuesta_satisfaccion.sql). Una sola encuesta por conversación.

Entrega: en vivo al cerrar (si el cliente tiene el chat abierto). Si no lo tenía abierto o recargó la página
sin responder, se le vuelve a mostrar al abrir el chat dentro de las 24 h siguientes (máximo 3 veces), salvo
que haya tocado "Ahora no".

Encuesta por SMS (opcional, migracion_encuesta_sms.sql + "encuesta_sms_activo" en el Editor): si el cliente no
llegó a ver la encuesta (salió del chat antes del cierre), a los N minutos se le manda un SMS con un enlace
corto (?es=<código>) que abre el widget SOLO con la encuesta, sin iniciar una conversación nueva. Un SMS de
encuesta por celular cada 7 días, solo de 8:00 a 20:59 y con tope diario. Ver pendientes_sms().

Las respuestas se guardan como CÓDIGOS estables; los textos viven aquí (PREGUNTAS). Si un día se cambia la
redacción de una opción, los datos históricos siguen siendo comparables.
"""

import io
import logging
import re
import secrets
from datetime import datetime, timedelta, timezone

from db import get_connection, release_connection

logger = logging.getLogger(__name__)
LIMA_TZ = timezone(timedelta(hours=-5))
VIGENCIA_HORAS = 24          # hasta cuándo se le vuelve a mostrar una encuesta sin responder
MAX_ENTREGAS = 3
MAX_TEXTO_OTRA = 300
MAX_COMENTARIO = 1000
SMS_VIGENCIA_HORAS = 72      # el enlace del SMS sirve 3 días
SMS_VENTANA_HORAS = 12       # no se manda SMS por encuestas de cierres más antiguos que esto
SMS_UNO_CADA_DIAS = 7        # máximo 1 SMS de encuesta por celular cada 7 días

SATISFACCION = [  # (valor, emoji, texto)
    (1, "😞", "Muy insatisfecho"), (2, "🙁", "Insatisfecho"), (3, "😐", "Neutral"),
    (4, "🙂", "Satisfecho"), (5, "🤩", "Muy satisfecho"),
]
FACILIDAD = [  # (valor, texto) — en el orden en que se muestran; 5 = más fácil
    (5, "Muy fácil"), (4, "Fácil"), (3, "Ni fácil ni difícil"), (2, "Difícil"), (1, "Muy difícil"),
]
DIFICULTADES = [  # (código, texto)
    ("ninguna", "No, ninguna"),
    ("encontrar_ingreso", "Encontrar dónde ingresar"),
    ("formulario", "Completar el formulario"),
    ("espera", "Esperar la atención"),
    ("retomar", "Continuar o retomar mi conversación"),
    ("otra", "Otra dificultad"),
]
MEJORAS = [  # (código, texto)
    ("acceso", "Encontrar más fácilmente el acceso al webchat"),
    ("rapidez", "Recibir una atención más rápida"),
    ("estado_solicitudes", "Poder consultar el estado de mis solicitudes"),
    ("adjuntar", "Adjuntar documentos o imágenes fácilmente"),
    ("notificaciones", "Recibir notificaciones sobre mis solicitudes"),
    ("retomar", "Poder retomar conversaciones anteriores"),
    ("autoservicio", "Contar con más opciones de autoservicio"),
    ("ninguna", "No necesito mejoras por ahora"),
    ("otra", "Otra mejora"),
]
_TXT_DIFICULTAD = dict(DIFICULTADES)
_TXT_MEJORA = dict(MEJORAS)
_TXT_FACILIDAD = dict(FACILIDAD)
_TXT_SATISFACCION = {v: t for v, _, t in SATISFACCION}

PREGUNTAS = {
    "titulo": "Tu opinión nos ayuda a mejorar",
    "subtitulo": "Encuesta de experiencia · Menos de 1 minuto",
    "intro": "En Sanitas queremos brindarte una atención cada vez más sencilla. "
             "Cuéntanos cómo te fue utilizando nuestro canal digital.",
    "satisfaccion": {"texto": "1. ¿Qué tan satisfecho estás con la atención recibida por nuestro webchat?",
                     "opciones": [{"valor": v, "emoji": e, "texto": t} for v, e, t in SATISFACCION]},
    "facilidad": {"texto": "2. ¿Qué tan fácil te resultó utilizar el webchat para realizar tu consulta?",
                  "opciones": [{"valor": v, "texto": t} for v, t in FACILIDAD]},
    "dificultad": {"texto": "3. ¿Tuviste alguna dificultad al utilizar el webchat?",
                   "opciones": [{"valor": c, "texto": t} for c, t in DIFICULTADES]},
    "mejoras": {"texto": "4. ¿Qué te gustaría que mejoráramos o incorporáramos en nuestro canal digital?",
                "ayuda": "Puedes elegir más de una opción.",
                "opciones": [{"valor": c, "texto": t} for c, t in MEJORAS]},
    "comentario": {"texto": "5. Si deseas, cuéntanos qué podríamos hacer para mejorar tu experiencia.",
                   "max": MAX_COMENTARIO},
    "max_otra": MAX_TEXTO_OTRA,
}

_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


def _limpiar_texto(valor, maximo: int) -> str | None:
    if not isinstance(valor, str):
        return None
    t = _CONTROL.sub("", valor).strip()
    t = re.sub(r"[ \t]+", " ", t)
    return t[:maximo] or None


def _payload(encuesta_id: int, anterior: bool) -> dict:
    return {"encuesta_id": encuesta_id, "anterior": anterior, "preguntas": PREGUNTAS}


# ---------------------------------------------------------------------------
# Envío
# ---------------------------------------------------------------------------
def crear_si_corresponde(conversacion_id: int) -> dict | None:
    """Crea la encuesta de esta conversación si aplica. Devuelve el payload para el widget o None.
    Si la migración aún no se corrió, no hace nada (no debe romper el cierre)."""
    conn = get_connection()
    try:
        with conn.cursor() as cur:
            cur.execute(
                """INSERT INTO encuestas_satisfaccion (conversacion_id, usuario_id, campana_id, asesor_id, enviada_en)
                   SELECT c.id, c.usuario_id, c.campana_id, c.asesor_id, %s
                   FROM conversaciones c JOIN campanas cam ON cam.id = c.campana_id
                   WHERE c.id = %s AND c.asesor_id IS NOT NULL AND cam.encuesta_satisfaccion = TRUE
                   ON CONFLICT (conversacion_id) DO NOTHING
                   RETURNING id""",
                (datetime.now(LIMA_TZ), conversacion_id),
            )
            fila = cur.fetchone()
        conn.commit()
        return _payload(fila[0], False) if fila else None
    except Exception:  # noqa: BLE001
        conn.rollback()
        logger.warning("No se pudo crear la encuesta de la conversación %s (¿falta la migración?)",
                       conversacion_id, exc_info=True)
        return None
    finally:
        release_connection(conn)


def marcar_entregada(encuesta_id: int):
    conn = get_connection()
    try:
        with conn.cursor() as cur:
            cur.execute("UPDATE encuestas_satisfaccion SET entregas = entregas + 1 WHERE id = %s", (encuesta_id,))
        conn.commit()
    except Exception:  # noqa: BLE001
        conn.rollback()
    finally:
        release_connection(conn)


def pendiente_para_token(sesion_token: str) -> dict | None:
    """La última encuesta sin responder de este visitante (últimas 24 h, máx. 3 veces). La marca como
    entregada en la misma operación, así dos pestañas no la cuentan dos veces."""
    conn = get_connection()
    try:
        with conn.cursor() as cur:
            cur.execute(
                """UPDATE encuestas_satisfaccion SET entregas = entregas + 1
                   WHERE id = (
                       SELECT e.id FROM encuestas_satisfaccion e
                       JOIN usuarios_webchat u ON u.id = e.usuario_id
                       WHERE u.sesion_token = %s AND e.respondida_en IS NULL AND e.omitida_en IS NULL
                         AND e.enviada_en > %s AND e.entregas < %s
                       ORDER BY e.enviada_en DESC LIMIT 1)
                   RETURNING id""",
                (sesion_token, datetime.now(LIMA_TZ) - timedelta(hours=VIGENCIA_HORAS), MAX_ENTREGAS),
            )
            fila = cur.fetchone()
        conn.commit()
        return _payload(fila[0], True) if fila else None
    except Exception:  # noqa: BLE001
        conn.rollback()
        return None
    finally:
        release_connection(conn)


# ---------------------------------------------------------------------------
# Respuesta
# ---------------------------------------------------------------------------
def _validar(datos: dict):
    """Devuelve (valores, None) o (None, motivo)."""
    try:
        satisfaccion = int(datos.get("satisfaccion"))
        facilidad = int(datos.get("facilidad"))
    except (TypeError, ValueError):
        return None, "Responde las preguntas 1 y 2."
    if satisfaccion not in _TXT_SATISFACCION:
        return None, "Responde la pregunta 1."
    if facilidad not in _TXT_FACILIDAD:
        return None, "Responde la pregunta 2."
    dificultad = datos.get("dificultad")
    if dificultad not in _TXT_DIFICULTAD:
        return None, "Responde la pregunta 3."
    mejoras = datos.get("mejoras")
    if not isinstance(mejoras, list) or not mejoras or any(m not in _TXT_MEJORA for m in mejoras):
        return None, "Elige al menos una opción en la pregunta 4."
    mejoras = [c for c, _ in MEJORAS if c in set(mejoras)]   # sin duplicados y en orden fijo
    if "ninguna" in mejoras and len(mejoras) > 1:
        mejoras = [m for m in mejoras if m != "ninguna"]
    return {
        "satisfaccion": satisfaccion,
        "facilidad": facilidad,
        "dificultad": dificultad,
        "dificultad_otra": _limpiar_texto(datos.get("dificultad_otra"), MAX_TEXTO_OTRA) if dificultad == "otra" else None,
        "mejoras": mejoras,
        "mejora_otra": _limpiar_texto(datos.get("mejora_otra"), MAX_TEXTO_OTRA) if "otra" in mejoras else None,
        "comentario": _limpiar_texto(datos.get("comentario"), MAX_COMENTARIO),
    }, None


def guardar_respuesta(encuesta_id, sesion_token: str, datos: dict) -> dict:
    """Guarda las respuestas. Solo el mismo visitante (sesion_token) puede responder su encuesta, y una vez."""
    if not str(encuesta_id or "").isdigit():
        return {"ok": False, "motivo": "Encuesta no válida."}
    valores, motivo = _validar(datos or {})
    if motivo:
        return {"ok": False, "motivo": motivo}
    conn = get_connection()
    try:
        with conn.cursor() as cur:
            cur.execute(
                """UPDATE encuestas_satisfaccion e
                   SET respondida_en = %s, omitida_en = NULL, satisfaccion = %s, facilidad = %s, dificultad = %s,
                       dificultad_otra = %s, mejoras = %s, mejora_otra = %s, comentario = %s
                   FROM usuarios_webchat u
                   WHERE e.id = %s AND u.id = e.usuario_id AND u.sesion_token = %s AND e.respondida_en IS NULL
                   RETURNING e.id""",
                (datetime.now(LIMA_TZ), valores["satisfaccion"], valores["facilidad"], valores["dificultad"],
                 valores["dificultad_otra"], valores["mejoras"], valores["mejora_otra"], valores["comentario"],
                 int(encuesta_id), sesion_token),
            )
            ok = cur.fetchone() is not None
            ya_respondida = False
            if not ok:
                cur.execute("""SELECT 1 FROM encuestas_satisfaccion e JOIN usuarios_webchat u ON u.id = e.usuario_id
                               WHERE e.id = %s AND u.sesion_token = %s AND e.respondida_en IS NOT NULL""",
                            (int(encuesta_id), sesion_token))
                ya_respondida = cur.fetchone() is not None
        conn.commit()
    except Exception:  # noqa: BLE001
        conn.rollback()
        logger.exception("No se pudo guardar la encuesta %s", encuesta_id)
        return {"ok": False, "motivo": "No pudimos guardar tus respuestas. Intenta de nuevo."}
    finally:
        release_connection(conn)
    if ok:
        return {"ok": True}
    if ya_respondida:   # p. ej. la respondió en otra pestaña: para el cliente es lo mismo que "gracias"
        return {"ok": False, "ya_respondida": True, "motivo": "Esta encuesta ya fue respondida."}
    return {"ok": False, "motivo": "Esta encuesta ya no está disponible."}


def omitir(encuesta_id, sesion_token: str) -> dict:
    if not str(encuesta_id or "").isdigit():
        return {"ok": False}
    conn = get_connection()
    try:
        with conn.cursor() as cur:
            cur.execute(
                """UPDATE encuestas_satisfaccion e SET omitida_en = %s
                   FROM usuarios_webchat u
                   WHERE e.id = %s AND u.id = e.usuario_id AND u.sesion_token = %s
                     AND e.respondida_en IS NULL AND e.omitida_en IS NULL""",
                (datetime.now(LIMA_TZ), int(encuesta_id), sesion_token),
            )
        conn.commit()
        return {"ok": True}
    except Exception:  # noqa: BLE001
        conn.rollback()
        return {"ok": False}
    finally:
        release_connection(conn)


# ---------------------------------------------------------------------------
# Encuesta por SMS (para quien no llegó a verla en el chat)
# ---------------------------------------------------------------------------
_aviso_migracion_sms = False


def _falta_migracion_sms(e: Exception) -> bool:
    global _aviso_migracion_sms
    if "sms_codigo" in str(e) or "sms_enviado_en" in str(e) or "sms_telefono" in str(e):
        if not _aviso_migracion_sms:
            logger.warning("encuesta: falta ejecutar migracion_encuesta_sms.sql (no se envían encuestas por SMS)")
            _aviso_migracion_sms = True
        return True
    return False


_RE_CELULAR = re.compile(r"^9\d{8}$")


def _celular_valido(t) -> str | None:
    d = re.sub(r"\D", "", str(t or ""))
    if d.startswith("51") and len(d) == 11:
        d = d[2:]
    return d if _RE_CELULAR.match(d) else None


def pendientes_sms(espera_min: int, campanas_codigos: list[str] | None, tope_diario: int,
                   ahora: datetime | None = None, limite: int = 20) -> list[dict]:
    """Encuestas a las que corresponde mandarles el SMS ahora. Reserva cada una (pone sms_codigo y
    sms_enviado_en) para no mandarla dos veces: [{"encuesta_id", "telefono", "codigo"}].
    Después de enviar, llamar a resultado_sms(encuesta_id, ok)."""
    ahora = ahora or datetime.now(LIMA_TZ)
    conn = get_connection()
    try:
        with conn.cursor() as cur:
            try:
                cur.execute("SELECT COUNT(*) FROM encuestas_satisfaccion WHERE sms_enviado_en >= %s",
                            (ahora.replace(hour=0, minute=0, second=0, microsecond=0),))
            except Exception as e:  # noqa: BLE001
                conn.rollback()
                if not _falta_migracion_sms(e):
                    logger.exception("encuesta: no se pudo revisar las encuestas por SMS")
                return []
            ya_hoy = cur.fetchone()[0]
            cupo = min(limite, max(0, tope_diario - ya_hoy))
            if cupo <= 0:
                conn.rollback()
                return []
            filtro_cam, params_cam = "", []
            if campanas_codigos:
                filtro_cam, params_cam = " AND cam.codigo = ANY(%s)", [list(campanas_codigos)]
            cur.execute(
                f"""SELECT e.id, COALESCE(NULLIF(u.telefono, ''), c.variables_sesion->>'CELULAR')
                    FROM encuestas_satisfaccion e
                    JOIN conversaciones c ON c.id = e.conversacion_id
                    JOIN usuarios_webchat u ON u.id = e.usuario_id
                    JOIN campanas cam ON cam.id = e.campana_id
                    LEFT JOIN etiquetas_cierre et ON et.id = c.etiqueta_cierre_id
                    WHERE e.respondida_en IS NULL AND e.omitida_en IS NULL AND e.sms_enviado_en IS NULL
                      AND e.entregas = 0
                      AND e.enviada_en <= %s AND e.enviada_en > %s
                      AND COALESCE(et.nombre, '') NOT ILIKE '%%prueba%%'
                      -- si volvió a escribir y está en cola o con un asesor, no se le interrumpe
                      AND NOT EXISTS (SELECT 1 FROM conversaciones c2 WHERE c2.usuario_id = e.usuario_id
                                        AND c2.estado IN ('en_cola', 'asignada')){filtro_cam}
                    ORDER BY e.enviada_en
                    LIMIT 200""",
                [ahora - timedelta(minutes=espera_min), ahora - timedelta(hours=SMS_VENTANA_HORAS)] + params_cam,
            )
            candidatos = [(i, _celular_valido(t)) for i, t in cur.fetchall()]
            candidatos = [(i, t) for i, t in candidatos if t]
            if not candidatos:
                conn.rollback()
                return []
            cur.execute("""SELECT DISTINCT sms_telefono FROM encuestas_satisfaccion
                           WHERE sms_telefono = ANY(%s) AND sms_enviado_en > %s""",
                        ([t for _, t in candidatos], ahora - timedelta(days=SMS_UNO_CADA_DIAS)))
            recientes = {r[0] for r in cur.fetchall()}
            salida = []
            for i, t in candidatos:
                if len(salida) >= cupo:
                    break
                if t in recientes:
                    continue
                recientes.add(t)   # tampoco dos encuestas del mismo celular en esta vuelta
                codigo = secrets.token_urlsafe(8)   # 11 caracteres: corto para el SMS, imposible de adivinar
                cur.execute("""UPDATE encuestas_satisfaccion SET sms_codigo = %s, sms_enviado_en = %s, sms_telefono = %s
                               WHERE id = %s AND sms_enviado_en IS NULL RETURNING id""", (codigo, ahora, t, i))
                if cur.fetchone():
                    salida.append({"encuesta_id": i, "telefono": t, "codigo": codigo})
        conn.commit()
        return salida
    except Exception:  # noqa: BLE001
        conn.rollback()
        logger.exception("encuesta: no se pudo preparar las encuestas por SMS")
        return []
    finally:
        release_connection(conn)


def resultado_sms(encuesta_id: int, ok: bool):
    conn = get_connection()
    try:
        with conn.cursor() as cur:
            cur.execute("UPDATE encuestas_satisfaccion SET sms_ok = %s WHERE id = %s", (ok, encuesta_id))
        conn.commit()
    except Exception:  # noqa: BLE001
        conn.rollback()
    finally:
        release_connection(conn)


def _por_codigo(cur, codigo: str):
    """(id, respondida_en, omitida_en, vigente, sesion_token) o None."""
    if not isinstance(codigo, str) or not re.fullmatch(r"[A-Za-z0-9_-]{6,32}", codigo):
        return None
    cur.execute("""SELECT e.id, e.respondida_en, e.omitida_en, e.sms_enviado_en > %s, u.sesion_token
                   FROM encuestas_satisfaccion e JOIN usuarios_webchat u ON u.id = e.usuario_id
                   WHERE e.sms_codigo = %s""",
                (datetime.now(LIMA_TZ) - timedelta(hours=SMS_VIGENCIA_HORAS), codigo))
    return cur.fetchone()


def abrir_por_sms(codigo: str) -> dict:
    """Lo que ve el cliente al abrir el enlace del SMS."""
    conn = get_connection()
    try:
        with conn.cursor() as cur:
            fila = _por_codigo(cur, codigo)
            if not fila:
                conn.rollback()
                return {"ok": False, "estado": "no_existe"}
            eid, respondida, omitida, vigente, token = fila
            if respondida:
                conn.rollback()
                return {"ok": False, "estado": "respondida", "sesion_token": token}
            if not vigente:
                conn.rollback()
                return {"ok": False, "estado": "vencida", "sesion_token": token}
            # Abrir el enlace la "reactiva" aunque antes haya tocado "Ahora no" en el chat.
            cur.execute("UPDATE encuestas_satisfaccion SET entregas = entregas + 1, omitida_en = NULL WHERE id = %s", (eid,))
        conn.commit()
        return {"ok": True, "payload": {**_payload(eid, True), "codigo_sms": codigo}, "sesion_token": token}
    except Exception:  # noqa: BLE001
        conn.rollback()
        logger.warning("encuesta: no se pudo abrir la encuesta por SMS", exc_info=True)
        return {"ok": False, "estado": "error"}
    finally:
        release_connection(conn)


def responder_por_sms(codigo: str, datos: dict) -> dict:
    valores, motivo = _validar(datos or {})
    if motivo:
        return {"ok": False, "motivo": motivo}
    conn = get_connection()
    try:
        with conn.cursor() as cur:
            fila = _por_codigo(cur, codigo)
            if not fila:
                conn.rollback()
                return {"ok": False, "motivo": "Esta encuesta ya no está disponible."}
            eid, respondida, _omitida, vigente, _token = fila
            if respondida:
                conn.rollback()
                return {"ok": False, "ya_respondida": True, "motivo": "Esta encuesta ya fue respondida."}
            if not vigente:
                conn.rollback()
                return {"ok": False, "motivo": "Esta encuesta ya venció. ¡Gracias de todas formas!"}
            cur.execute(
                """UPDATE encuestas_satisfaccion
                   SET respondida_en = %s, omitida_en = NULL, satisfaccion = %s, facilidad = %s, dificultad = %s,
                       dificultad_otra = %s, mejoras = %s, mejora_otra = %s, comentario = %s, respondida_por_sms = TRUE
                   WHERE id = %s AND respondida_en IS NULL""",
                (datetime.now(LIMA_TZ), valores["satisfaccion"], valores["facilidad"], valores["dificultad"],
                 valores["dificultad_otra"], valores["mejoras"], valores["mejora_otra"], valores["comentario"], eid),
            )
        conn.commit()
        return {"ok": True, "encuesta_id": eid}
    except Exception:  # noqa: BLE001
        conn.rollback()
        logger.exception("No se pudo guardar la encuesta por SMS")
        return {"ok": False, "motivo": "No pudimos guardar tus respuestas. Intenta de nuevo."}
    finally:
        release_connection(conn)


def omitir_por_sms(codigo: str) -> dict:
    conn = get_connection()
    try:
        with conn.cursor() as cur:
            fila = _por_codigo(cur, codigo)
            if fila and not fila[1]:
                cur.execute("UPDATE encuestas_satisfaccion SET omitida_en = %s WHERE id = %s AND respondida_en IS NULL",
                            (datetime.now(LIMA_TZ), fila[0]))
        conn.commit()
        return {"ok": True}
    except Exception:  # noqa: BLE001
        conn.rollback()
        return {"ok": False}
    finally:
        release_connection(conn)


# ---------------------------------------------------------------------------
# Resultados (dashboard y Excel)
# ---------------------------------------------------------------------------
def _filtro_campanas(campana_ids) -> tuple[str, list]:
    if campana_ids is None:
        return "", []
    return " AND e.campana_id = ANY(%s)", [list(int(i) for i in campana_ids)]


def _pct(a, b):
    return round(a * 100.0 / b, 1) if b else None


def estadisticas(f_desde: datetime, f_hasta: datetime, campana_ids=None) -> dict:
    """Resultados de las encuestas ENVIADAS en el período (hora de Lima). campana_ids None = todas."""
    extra, params_extra = _filtro_campanas(campana_ids)
    base = [f_desde, f_hasta] + params_extra
    donde = "e.enviada_en >= %s AND e.enviada_en < %s" + extra
    conn = get_connection()
    try:
        with conn.cursor() as cur:
            try:
                cur.execute(f"SELECT COUNT(*), COUNT(e.respondida_en), COUNT(e.omitida_en) FROM encuestas_satisfaccion e WHERE {donde}", base)
            except Exception:  # noqa: BLE001
                conn.rollback()
                return {"disponible": False}
            enviadas, respondidas, omitidas = cur.fetchone()
            r = donde + " AND e.respondida_en IS NOT NULL"

            cur.execute(f"SELECT e.satisfaccion, COUNT(*) FROM encuestas_satisfaccion e WHERE {r} GROUP BY 1", base)
            dist_sat = dict(cur.fetchall())
            cur.execute(f"SELECT e.facilidad, COUNT(*) FROM encuestas_satisfaccion e WHERE {r} GROUP BY 1", base)
            dist_fac = dict(cur.fetchall())

            # Dificultades, con la satisfacción promedio de quienes la marcaron (¿cuánto pesa en la experiencia?)
            cur.execute(f"""SELECT e.dificultad, COUNT(*), AVG(e.satisfaccion)
                            FROM encuestas_satisfaccion e WHERE {r} GROUP BY 1""", base)
            dif = {c: (n, float(p) if p is not None else None) for c, n, p in cur.fetchall()}

            cur.execute(f"""SELECT m, COUNT(*), AVG(e.satisfaccion)
                            FROM encuestas_satisfaccion e, unnest(e.mejoras) AS m WHERE {r} GROUP BY 1""", base)
            mej = {c: (n, float(p) if p is not None else None) for c, n, p in cur.fetchall()}

            cur.execute(f"""SELECT COALESCE(cam.nombre, 'Sin campaña'), COUNT(*), COUNT(e.respondida_en),
                                   AVG(e.satisfaccion), COUNT(*) FILTER (WHERE e.satisfaccion >= 4),
                                   COUNT(*) FILTER (WHERE e.facilidad >= 4)
                            FROM encuestas_satisfaccion e LEFT JOIN campanas cam ON cam.id = e.campana_id
                            WHERE {donde} GROUP BY 1 ORDER BY 2 DESC""", base)
            por_campana = [{"campana": n, "enviadas": env, "respondidas": resp,
                            "promedio": round(float(p), 2) if p is not None else None,
                            "csat": _pct(s, resp), "ces": _pct(f, resp)}
                           for n, env, resp, p, s, f in cur.fetchall()]

            cur.execute(f"""SELECT COALESCE(a.nombre_display, 'Sin asesor'), COUNT(*), COUNT(e.respondida_en),
                                   AVG(e.satisfaccion), COUNT(*) FILTER (WHERE e.satisfaccion >= 4),
                                   COUNT(*) FILTER (WHERE e.satisfaccion <= 2)
                            FROM encuestas_satisfaccion e LEFT JOIN asesores a ON a.id = e.asesor_id
                            WHERE {donde} GROUP BY 1 ORDER BY 3 DESC, 1""", base)
            por_asesor = [{"asesor": n, "enviadas": env, "respondidas": resp,
                           "promedio": round(float(p), 2) if p is not None else None,
                           "csat": _pct(s, resp), "bajas": b}
                          for n, env, resp, p, s, b in cur.fetchall()]

            cur.execute(f"""SELECT (e.enviada_en AT TIME ZONE 'America/Lima')::date, COUNT(*), COUNT(e.respondida_en),
                                   COUNT(*) FILTER (WHERE e.satisfaccion >= 4), AVG(e.satisfaccion)
                            FROM encuestas_satisfaccion e WHERE {donde} GROUP BY 1 ORDER BY 1""", base)
            por_dia = [{"dia": d.isoformat(), "enviadas": env, "respondidas": resp, "csat": _pct(s, resp),
                        "promedio": round(float(p), 2) if p is not None else None}
                       for d, env, resp, s, p in cur.fetchall()]

            # Encuesta por SMS (si ya se corrió migracion_encuesta_sms.sql)
            sms = None
            try:
                cur.execute("SAVEPOINT sp_sms")
                cur.execute(f"""SELECT COUNT(*) FILTER (WHERE e.sms_enviado_en IS NOT NULL AND e.sms_ok IS NOT FALSE),
                                       COUNT(*) FILTER (WHERE e.sms_ok IS FALSE),
                                       COUNT(*) FILTER (WHERE e.respondida_por_sms)
                                FROM encuestas_satisfaccion e WHERE {donde}""", base)
                env_sms, fall_sms, resp_sms = cur.fetchone()
                sms = {"enviados": env_sms, "fallidos": fall_sms, "respondidas": resp_sms,
                       "tasa_respuesta": _pct(resp_sms, env_sms)}
                cur.execute("RELEASE SAVEPOINT sp_sms")
            except Exception:  # noqa: BLE001
                cur.execute("ROLLBACK TO SAVEPOINT sp_sms")

            cur.execute(f"""SELECT e.respondida_en, e.conversacion_id, COALESCE(cam.nombre, ''), COALESCE(a.nombre_display, ''),
                                   e.satisfaccion, e.facilidad, e.dificultad, e.dificultad_otra, e.mejoras, e.mejora_otra, e.comentario
                            FROM encuestas_satisfaccion e
                            LEFT JOIN campanas cam ON cam.id = e.campana_id LEFT JOIN asesores a ON a.id = e.asesor_id
                            WHERE {r} AND (e.comentario IS NOT NULL OR e.dificultad_otra IS NOT NULL OR e.mejora_otra IS NOT NULL)
                            ORDER BY e.respondida_en DESC LIMIT 100""", base)
            textos = [{"fecha": f.astimezone(LIMA_TZ).strftime("%d/%m/%Y %H:%M"), "conversacion_id": cid, "campana": cam,
                       "asesor": ase, "satisfaccion": s, "facilidad": _TXT_FACILIDAD.get(fa),
                       "dificultad": _TXT_DIFICULTAD.get(d), "dificultad_otra": do,
                       "mejoras": [_TXT_MEJORA.get(m, m) for m in (ms or [])], "mejora_otra": mo, "comentario": co}
                      for f, cid, cam, ase, s, fa, d, do, ms, mo, co in cur.fetchall()]
        conn.rollback()
    finally:
        release_connection(conn)

    resp = respondidas
    prom_sat = (sum(k * v for k, v in dist_sat.items()) / resp) if resp else None
    prom_fac = (sum(k * v for k, v in dist_fac.items()) / resp) if resp else None
    satisfechos = dist_sat.get(4, 0) + dist_sat.get(5, 0)
    insatisfechos = dist_sat.get(1, 0) + dist_sat.get(2, 0)
    faciles = dist_fac.get(4, 0) + dist_fac.get(5, 0)
    dificiles = dist_fac.get(1, 0) + dist_fac.get(2, 0)

    dificultades = [{"codigo": c, "texto": t, "cantidad": dif.get(c, (0, None))[0], "pct": _pct(dif.get(c, (0, None))[0], resp),
                     "satisfaccion_prom": round(dif[c][1], 2) if c in dif and dif[c][1] is not None else None}
                    for c, t in DIFICULTADES]
    mejoras = sorted(
        [{"codigo": c, "texto": t, "cantidad": mej.get(c, (0, None))[0], "pct": _pct(mej.get(c, (0, None))[0], resp),
          "satisfaccion_prom": round(mej[c][1], 2) if c in mej and mej[c][1] is not None else None}
         for c, t in MEJORAS],
        key=lambda x: (x["codigo"] in ("ninguna", "otra"), -x["cantidad"]))

    return {
        "disponible": True,
        "enviadas": enviadas, "respondidas": resp, "omitidas": omitidas,
        "tasa_respuesta": _pct(resp, enviadas),
        "sms": sms,
        "satisfaccion": {"promedio": round(prom_sat, 2) if prom_sat is not None else None,
                         "csat": _pct(satisfechos, resp), "insatisfechos_pct": _pct(insatisfechos, resp),
                         "distribucion": [{"valor": v, "emoji": e, "texto": t, "cantidad": dist_sat.get(v, 0),
                                           "pct": _pct(dist_sat.get(v, 0), resp)} for v, e, t in SATISFACCION]},
        "facilidad": {"promedio": round(prom_fac, 2) if prom_fac is not None else None,
                      "ces": _pct(faciles, resp), "dificil_pct": _pct(dificiles, resp),
                      "distribucion": [{"valor": v, "texto": t, "cantidad": dist_fac.get(v, 0),
                                        "pct": _pct(dist_fac.get(v, 0), resp)} for v, t in FACILIDAD]},
        "dificultades": dificultades,
        "mejoras": mejoras,
        "por_campana": por_campana,
        "por_asesor": por_asesor,
        "por_dia": por_dia,
        "textos": textos,
        "hallazgos": _hallazgos(resp, prom_sat, _pct(satisfechos, resp), _pct(faciles, resp),
                                dificultades, mejoras, por_campana, _pct(resp, enviadas)),
    }


def _n(v, dec=1) -> str:
    return f"{v:.{dec}f}".replace(".", ",")


def _hallazgos(resp, prom_sat, csat, ces, dificultades, mejoras, por_campana, tasa) -> list:
    """Frases cortas, calculadas, para orientar las mejoras del canal. Sin opiniones: solo lo que dicen los datos."""
    if not resp:
        return []
    out = []
    MIN = 5   # no sacar conclusiones con muy pocas respuestas
    if resp < MIN:
        out.append(f"Todavía hay pocas respuestas ({resp}): espera a tener más antes de sacar conclusiones.")
    con_dif = [d for d in dificultades if d["codigo"] != "ninguna" and d["cantidad"]]
    sin_dif = next((d for d in dificultades if d["codigo"] == "ninguna"), None)
    if sin_dif and sin_dif["pct"] is not None:
        out.append(f"El {100 - sin_dif['pct']:.0f}% tuvo alguna dificultad al usar el webchat.")
    if con_dif:
        top = max(con_dif, key=lambda d: d["cantidad"])
        out.append(f"La dificultad más mencionada es «{top['texto']}» ({top['pct']:.0f}% de las respuestas).")
        if prom_sat is not None:
            peor = min((d for d in con_dif if d["cantidad"] >= MIN and d["satisfaccion_prom"] is not None),
                       key=lambda d: d["satisfaccion_prom"], default=None)
            if peor and peor["satisfaccion_prom"] < prom_sat - 0.3:
                out.append(f"Quienes marcaron «{peor['texto']}» califican la atención con {_n(peor['satisfaccion_prom'])} de 5, "
                           f"frente a {_n(prom_sat)} en general: es la que más baja la satisfacción.")
    pedidas = [m for m in mejoras if m["codigo"] not in ("ninguna", "otra") and m["cantidad"]]
    if pedidas:
        top = pedidas[:3]
        out.append("Lo más pedido: " + "; ".join(f"«{m['texto']}» ({m['pct']:.0f}%)" for m in top) + ".")
    cams = [c for c in por_campana if (c["respondidas"] or 0) >= MIN and c["csat"] is not None]
    if len(cams) >= 2:
        peor = min(cams, key=lambda c: c["csat"])
        mejor = max(cams, key=lambda c: c["csat"])
        if mejor["csat"] - peor["csat"] >= 10:
            out.append(f"«{peor['campana']}» tiene la satisfacción más baja ({peor['csat']:.0f}% satisfechos) "
                       f"frente a «{mejor['campana']}» ({mejor['csat']:.0f}%).")
    if ces is not None and csat is not None and ces < csat - 15:
        out.append(f"La atención se valora mejor ({csat:.0f}% satisfechos) que la facilidad de uso del canal "
                   f"({ces:.0f}% lo encontró fácil): el problema está más en el canal que en el asesor.")
    if tasa is not None and tasa < 20:
        out.append(f"Solo responde el {tasa:.0f}% de quienes reciben la encuesta: los resultados pueden no representar a todos.")
    return out


def exportar_xlsx(f_desde: datetime, f_hasta: datetime, campana_ids, etiqueta: str) -> bytes:
    """Excel con una fila por encuesta respondida (con los textos de las opciones) y una hoja de resumen."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    extra, params_extra = _filtro_campanas(campana_ids)
    conn = get_connection()
    try:
        with conn.cursor() as cur:
            cur.execute(
                f"""SELECT e.respondida_en, e.conversacion_id, COALESCE(cam.nombre, ''), COALESCE(a.nombre_display, ''),
                           e.satisfaccion, e.facilidad, e.dificultad, e.dificultad_otra, e.mejoras, e.mejora_otra, e.comentario
                    FROM encuestas_satisfaccion e
                    LEFT JOIN campanas cam ON cam.id = e.campana_id LEFT JOIN asesores a ON a.id = e.asesor_id
                    WHERE e.enviada_en >= %s AND e.enviada_en < %s AND e.respondida_en IS NOT NULL{extra}
                    ORDER BY e.respondida_en""",
                [f_desde, f_hasta] + params_extra,
            )
            filas = cur.fetchall()
        conn.rollback()
    finally:
        release_connection(conn)
    est = estadisticas(f_desde, f_hasta, campana_ids)

    wb = Workbook()
    titulo = Font(bold=True, size=13, color="0A4C8C")
    cab = Font(bold=True, color="FFFFFF")
    relleno = PatternFill("solid", fgColor="0A4C8C")
    limpio = lambda v: _CONTROL.sub("", v) if isinstance(v, str) else v  # noqa: E731

    ws = wb.active
    ws.title = "Resumen"
    ws.append([f"Encuesta de satisfacción — {etiqueta}"]); ws["A1"].font = titulo
    ws.append([])
    s, f = est.get("satisfaccion", {}), est.get("facilidad", {})
    for et, v in [("Encuestas enviadas", est.get("enviadas")), ("Respondidas", est.get("respondidas")),
                  ("Tasa de respuesta (%)", est.get("tasa_respuesta")), ("Satisfacción promedio (1 a 5)", s.get("promedio")),
                  ("Satisfechos — 4 y 5 (%)", s.get("csat")), ("Insatisfechos — 1 y 2 (%)", s.get("insatisfechos_pct")),
                  ("Lo encontró fácil o muy fácil (%)", f.get("ces")), ("Lo encontró difícil o muy difícil (%)", f.get("dificil_pct"))]:
        ws.append([et, v])
    ws.append([])
    ws.append(["Hallazgos"]); ws.cell(ws.max_row, 1).font = Font(bold=True)
    for h in est.get("hallazgos", []):
        ws.append([h])
    ws.append([])
    ws.append(["Dificultades", "Cantidad", "%", "Satisfacción promedio"])
    for c in ws[ws.max_row]: c.font = cab; c.fill = relleno
    for d in est.get("dificultades", []):
        ws.append([d["texto"], d["cantidad"], d["pct"], d["satisfaccion_prom"]])
    ws.append([])
    ws.append(["Mejoras pedidas", "Cantidad", "%", "Satisfacción promedio"])
    for c in ws[ws.max_row]: c.font = cab; c.fill = relleno
    for m in est.get("mejoras", []):
        ws.append([m["texto"], m["cantidad"], m["pct"], m["satisfaccion_prom"]])
    ws.append([])
    ws.append(["Campaña", "Enviadas", "Respondidas", "Satisfacción promedio", "Satisfechos (%)", "Fácil (%)"])
    for c in ws[ws.max_row]: c.font = cab; c.fill = relleno
    for c in est.get("por_campana", []):
        ws.append([c["campana"], c["enviadas"], c["respondidas"], c["promedio"], c["csat"], c["ces"]])
    ws.column_dimensions["A"].width = 60
    for col in "BCDEF":
        ws.column_dimensions[col].width = 18

    ws2 = wb.create_sheet("Respuestas")
    encab = ["Fecha respuesta", "Conversación", "Campaña", "Asesor", "Satisfacción (1-5)", "Satisfacción",
             "Facilidad", "Dificultad", "Otra dificultad", "Mejoras", "Otra mejora", "Comentario"]
    ws2.append(encab)
    for c in ws2[1]: c.font = cab; c.fill = relleno
    for fr, cid, cam, ase, sa, fa, di, do, ms, mo, co in filas:
        ws2.append([fr.astimezone(LIMA_TZ).replace(tzinfo=None), cid, limpio(cam), limpio(ase), sa, _TXT_SATISFACCION.get(sa),
                    _TXT_FACILIDAD.get(fa), _TXT_DIFICULTAD.get(di, di), limpio(do),
                    "; ".join(_TXT_MEJORA.get(m, m) for m in (ms or [])), limpio(mo), limpio(co)])
        ws2.cell(ws2.max_row, 1).number_format = "dd/mm/yyyy hh:mm"
    for i, w in enumerate([17, 13, 26, 20, 12, 18, 20, 32, 30, 60, 30, 70], start=1):
        ws2.column_dimensions[ws2.cell(1, i).column_letter].width = w
    for fila in ws2.iter_rows(min_row=2):
        for c in fila:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = ws2.dimensions

    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()
